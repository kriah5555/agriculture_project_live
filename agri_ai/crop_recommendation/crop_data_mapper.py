"""
agri_ai.crop_recommendation.crop_data_mapper
-----------------------------------------------
Reads World_Country_Wise_Crop_Database_UPDATED.xlsx and provides lookup of
the "Crop Guide" fields (duration, sowing season, water requirement, pests,
fertilizers, harvest time) by (state, crop name).

Ported from the AgroZone prototype (apps/engine/crop_data_mapper.py),
dropped the Django `settings.DATA_ROOT` dependency in favour of a path
relative to this module — same self-contained style as agri_ai.fertilizer.

Dataset format: one sheet per country, headers at row 6, data from row 8.
Additional-info columns: 21 Duration, 22 Sowing Season, 23 Water Requirement,
24 Common Pests, 25 Recommended Fertilizers, 26 Harvest Time.
"""
import os

import openpyxl

from .crop_utils import normalize_crop_name, extract_base_name, safe_get, INDIAN_TO_ENGLISH_CROP_NAMES

DATA_DIR  = os.path.join(os.path.dirname(__file__), 'data')
WORLD_XLSX = os.path.join(DATA_DIR, 'World_Country_Wise_Crop_Database_UPDATED.xlsx')

COL_CROP_NAME               = 2
COL_CROP_DURATION           = 21
COL_SOWING_SEASON           = 22
COL_WATER_REQ               = 23
COL_COMMON_PESTS            = 24
COL_RECOMMENDED_FERTILIZERS = 25
COL_HARVEST_TIME            = 26
DATA_START_ROW = 8

# India dataset only has one country sheet — every state maps to it.
_STATE_TO_COUNTRY_DEFAULT = "India"


class CropDataMapper:
    def __init__(self, excel_path=None):
        self.excel_path = excel_path or WORLD_XLSX
        self._index = {}
        self._fallback = {}
        self._loaded = False

    def _resolve_country(self, location_name):
        if not location_name:
            return _STATE_TO_COUNTRY_DEFAULT
        # All India states share the single "India" sheet.
        return _STATE_TO_COUNTRY_DEFAULT

    def load(self):
        if self._loaded:
            return
        if not os.path.exists(self.excel_path):
            self._loaded = True
            return

        wb = openpyxl.load_workbook(self.excel_path, read_only=True)
        for sheet_name in wb.sheetnames:
            if sheet_name in ("INDEX", "Soil_Methodology"):
                continue
            ws = wb[sheet_name]
            country_key = sheet_name.strip().lower()

            for row in ws.iter_rows(min_row=DATA_START_ROW, values_only=True):
                if not row or not row[0]:
                    continue
                crop_name_val = row[COL_CROP_NAME - 1] if len(row) >= COL_CROP_NAME else None
                if not crop_name_val:
                    continue
                crop_name = str(crop_name_val).strip()
                if not crop_name:
                    continue

                cat_val = str(row[0]).strip() if row[0] else ""
                if cat_val and (cat_val.isupper() or cat_val.startswith("  ")):
                    continue

                entry = {
                    "country": sheet_name,
                    "crop_name": crop_name,
                    "crop_duration": safe_get(row[COL_CROP_DURATION - 1] if len(row) >= COL_CROP_DURATION else None),
                    "best_sowing_season": safe_get(row[COL_SOWING_SEASON - 1] if len(row) >= COL_SOWING_SEASON else None),
                    "water_requirement": safe_get(row[COL_WATER_REQ - 1] if len(row) >= COL_WATER_REQ else None),
                    "common_pests": safe_get(row[COL_COMMON_PESTS - 1] if len(row) >= COL_COMMON_PESTS else None),
                    "recommended_fertilizers": safe_get(row[COL_RECOMMENDED_FERTILIZERS - 1] if len(row) >= COL_RECOMMENDED_FERTILIZERS else None),
                    "harvest_time": safe_get(row[COL_HARVEST_TIME - 1] if len(row) >= COL_HARVEST_TIME else None),
                }

                crop_norm  = normalize_crop_name(crop_name)
                crop_lower = crop_name.lower()

                self._index[(country_key, crop_norm)] = entry
                self._index[(country_key, crop_lower)] = entry
                if crop_norm not in self._fallback:
                    self._fallback[crop_norm] = entry
                if crop_lower not in self._fallback:
                    self._fallback[crop_lower] = entry

        wb.close()
        self._loaded = True

    def lookup(self, crop_name, location_name=None):
        self.load()
        if not crop_name:
            return None

        country = self._resolve_country(location_name)
        country_key = country.strip().lower() if country else None

        crop_norm  = normalize_crop_name(crop_name)
        crop_lower = crop_name.strip().lower()
        base_name  = extract_base_name(crop_name)
        base_norm  = normalize_crop_name(base_name)

        if country_key:
            for key in [(country_key, crop_norm), (country_key, crop_lower)]:
                if key in self._index:
                    return self._index[key]
            if base_norm:
                key = (country_key, base_norm)
                if key in self._index:
                    return self._index[key]

        for key in [crop_norm, crop_lower]:
            if key in self._fallback:
                return self._fallback[key]

        if base_norm and base_norm not in (crop_norm, crop_lower):
            if base_norm in self._fallback:
                return self._fallback[base_norm]

        local_lower = crop_norm.strip()
        if local_lower in INDIAN_TO_ENGLISH_CROP_NAMES:
            english_name  = INDIAN_TO_ENGLISH_CROP_NAMES[local_lower]
            english_norm  = normalize_crop_name(english_name)
            english_lower = english_name.lower()
            if country_key:
                for key in [(country_key, english_norm), (country_key, english_lower)]:
                    if key in self._index:
                        return self._index[key]
            for key in [english_norm, english_lower]:
                if key in self._fallback:
                    return self._fallback[key]

        for fb_key, fb_entry in self._fallback.items():
            if fb_key in crop_norm or crop_norm in fb_key:
                return fb_entry
            if base_norm and (fb_key in base_norm or base_norm in fb_key):
                return fb_entry

        return None


_mapper_instance = None


def get_mapper():
    global _mapper_instance
    if _mapper_instance is None:
        _mapper_instance = CropDataMapper()
    return _mapper_instance