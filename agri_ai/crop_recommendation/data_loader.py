"""
agri_ai.crop_recommendation.data_loader
----------------------------------------
Reads India_State_Wise_Crops_Final.xlsx directly (no DB models, no migrations —
same pattern as agri_ai.fertilizer.data_loader) and exposes, per state, the
list of crops with their pH/OC/EC/N/P/K/S/Fe/Mn/Cu/Zn/B/Ca/Mg tolerance
ranges + ideal temperature range.

Columns per state sheet:
  Category(0) | Crop Name(1) | Notes(2) | Zone(3) |
  OC(4) | N(5) | P2O5(6) | K(7) | S(8) |
  Fe(9) | Mn(10) | Cu(11) | Zn(12) | B(13) |
  pH(14) | EC(15) | Ca(16) | Mg(17) | Temp(18)
"""
import os
import re

from openpyxl import load_workbook

DATA_DIR   = os.path.join(os.path.dirname(__file__), 'data')
CROP_XLSX  = os.path.join(DATA_DIR, 'India_State_Wise_Crops_Final.xlsx')

SKIP_SHEETS = {"INDEX", "Soil_Methodology", "Sheet1", "Sheet2", "Sheet3", "Sheet4"}

_PARAM_COLS = [
    ('oc', 4), ('n', 5), ('p', 6), ('k', 7), ('s', 8),
    ('fe', 9), ('mn', 10), ('cu', 11), ('zn', 12), ('b', 13),
    ('ph', 14), ('ec', 15), ('ca', 16), ('mg', 17),
]

_cache = None  # {state_name: [crop_dict, ...]}


def _parse_range(value) -> tuple:
    """Parse a range string like '5.5-7.5' or '240-292' into (lo, hi) floats."""
    if not value:
        return 0.0, 0.0
    value = str(value).strip()
    m = re.match(r"^([\d.]+)\s*[-–]\s*([\d.]+)", value)
    if m:
        return float(m.group(1)), float(m.group(2))
    try:
        v = float(value)
        return v, v
    except ValueError:
        return 0.0, 0.0


def _load_all():
    global _cache
    if _cache is not None:
        return _cache

    wb = load_workbook(CROP_XLSX, read_only=True, data_only=True)
    data = {}

    for sheet_name in wb.sheetnames:
        if sheet_name in SKIP_SHEETS:
            continue
        ws = wb[sheet_name]

        notable_facts = []
        crops = []
        sheet_order = 0

        rows = list(ws.iter_rows(values_only=True))
        for row in rows:
            if row and row[0] and str(row[0]).startswith("★"):
                notable_facts.append(str(row[0]).lstrip("★").strip())

        for row in rows:
            if not row or not row[0] or not row[1]:
                continue
            cat = str(row[0]).strip()
            if cat.startswith("  ") or cat.startswith("★") or cat.startswith("🌾") or cat == "Category":
                continue

            crop_name = str(row[1]).strip()
            notes     = str(row[2]).strip() if row[2] else ""
            zone_name = str(row[3]).strip() if row[3] else ""
            temp      = str(row[18]).strip() if len(row) > 18 and row[18] else ""
            is_notable = any(crop_name.lower() in fact.lower() for fact in notable_facts)

            crop = {
                'category':   cat,
                'crop_name':  crop_name,
                'notes':      notes,
                'zone_name':  zone_name,
                'temp_range': temp,
                'is_notable': is_notable,
                'sheet_order': sheet_order,
            }
            for param, col_idx in _PARAM_COLS:
                lo, hi = _parse_range(row[col_idx])
                crop[f'{param}_lo'] = lo
                crop[f'{param}_hi'] = hi

            crops.append(crop)
            sheet_order += 1

        data[sheet_name] = crops

    wb.close()
    _cache = data
    return _cache


def get_states() -> list:
    """Alphabetically sorted list of India states available in the dataset."""
    return sorted(_load_all().keys())


def get_crops_for_state(state: str) -> list:
    """List of crop dicts (with soil-range fields) for a given state name."""
    return _load_all().get(state, [])