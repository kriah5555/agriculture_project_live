"""
TEMPORARY FEATURE — admin viewer for ChannelData (raw sensor channel readings
captured by devise_apis.views.add_soil_data_open via channel_data[<key>] params).

ChannelData rows auto-expire after ChannelData.RETENTION_DAYS (see agriapp/models.py) —
this whole viewer only exists to inspect/export that short-lived data while it's around.
Kept deliberately self-contained so it's easy to rip out later:
  - delete this file
  - delete agriapp/channel_data_urls.py
  - remove its include() line from agriapp/urls.py
  - remove the "Channel Data" link in templates/navbar.html
  - delete templates/agriapp/channel_data_devices.html and channel_data_list.html
(ChannelData the model / its retention purge in devise_apis/views.py are unrelated to
this viewer and should stay regardless of whether this admin UI is removed.)
"""
import openpyxl
from openpyxl.styles import Font, PatternFill
from django.db.models import Count, Max
from django.http import HttpResponse, JsonResponse
from django.shortcuts import render, get_object_or_404

from .models import Devise, ChannelData, CHANNEL_EXPORT_COLUMNS
from .views import admin_required


@admin_required
def channel_data_devices(request):
    devices = (
        Devise.objects
        .filter(channeldata__isnull=False)
        .annotate(channel_count=Count('channeldata'), latest_reading=Max('channeldata__created_at'))
        .distinct()
        .order_by('-latest_reading')
    )
    context = {
        'devices': devices,
        'active_page': 'channel-data',
    }
    return render(request, 'agriapp/channel_data_devices.html', context)


@admin_required
def channel_data_list(request, pk):
    device = get_object_or_404(Devise, pk=pk)
    readings = ChannelData.objects.filter(device=device).order_by('-created_at')
    context = {
        'device': device,
        'readings': readings,
        'active_page': 'channel-data',
    }
    return render(request, 'agriapp/channel_data_list.html', context)


@admin_required
def channel_data_detail(request, pk):
    reading = get_object_or_404(ChannelData, pk=pk)
    data = {label: getattr(reading, field) for field, label in CHANNEL_EXPORT_COLUMNS}
    data['Recorded At'] = reading.created_at.strftime('%d %b %Y %H:%M')
    return JsonResponse(data)


@admin_required
def channel_data_export(request, pk):
    device = get_object_or_404(Devise, pk=pk)
    readings = ChannelData.objects.filter(device=device).order_by('-created_at')

    wb = openpyxl.Workbook()
    ws = wb.active
    ws.title = "Channel Data"

    bold_font  = Font(bold=True, color="FFFFFF")
    fill_color = PatternFill(start_color="4F81BD", end_color="4F81BD", fill_type="solid")

    headers = ['Recorded At'] + [label for _, label in CHANNEL_EXPORT_COLUMNS]
    ws.append(headers)
    for col_num in range(1, len(headers) + 1):
        cell = ws.cell(row=1, column=col_num)
        cell.font = bold_font
        cell.fill = fill_color

    for reading in readings:
        row = [reading.created_at.replace(tzinfo=None)]
        row += [getattr(reading, field) for field, _ in CHANNEL_EXPORT_COLUMNS]
        ws.append(row)

    for col in ws.columns:
        max_length = max((len(str(cell.value)) for cell in col if cell.value is not None), default=0)
        ws.column_dimensions[col[0].column_letter].width = max_length + 2

    response = HttpResponse(content_type="application/vnd.openxmlformats-officedocument.spreadsheetml.sheet")
    response["Content-Disposition"] = f'attachment; filename="channel_data_{device.devise_id or device.pk}.xlsx"'
    wb.save(response)
    return response
